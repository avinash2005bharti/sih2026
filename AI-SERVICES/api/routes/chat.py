"""
Chat endpoints with LangGraph Agent execution, multi-model routing, memory enrichment,
dedicated OCR + Moondream visual understanding pipeline, and real-time SSE streaming.
"""

import os
import sys
from pathlib import Path

# Ensure AI-SERVICES root is in sys.path regardless of working directory
_ai_root = str(Path(__file__).resolve().parent.parent.parent)
if _ai_root not in sys.path:
    sys.path.insert(0, _ai_root)

import json
import uuid
import asyncio
from typing import Optional, Dict, Any, List
from fastapi import APIRouter, HTTPException, UploadFile, File
import mimetypes
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field, model_validator

from llm.ollama_client import ollama_client
from llm.model_router import model_router
from agents.langgraph_agent import get_langgraph_agent
from memory import memory_manager
from memory import central_context_builder
from memory import conversation_summarizer
from memory import stm_manager
from multimodal.multimodal_orchestrator import multimodal_orchestrator
from core.logging import logger
from core.config import settings
from orchestrator.task_classifier import task_classifier, system_prompt_for

router = APIRouter()


class ChatRequest(BaseModel):
    """Chat request payload."""
    message: str = Field(..., description="User message or prompt")
    model: str = Field(default="auto", description="Model name or 'auto' for dynamic routing")
    agent: Optional[str] = Field(default="auto", description="Agent slug or 'auto' for central routing")
    conversation_id: Optional[str] = Field(None, description="Conversation ID for context retention")
    user_id: Optional[str] = Field(None, description="User ID for RBAC and episodic memory recall")
    request_id: Optional[str] = Field(None, description="End-to-end correlation ID")
    task_type: Optional[str] = Field(default="auto", description="Task type classification hint")
    system_prompt: Optional[str] = Field(None, description="Custom system instructions")
    temperature: Optional[float] = Field(None, ge=0.0, le=2.0, description="Sampling temperature")
    max_tokens: Optional[int] = Field(None, ge=1, le=8192, description="Maximum tokens to generate")
    images: Optional[List[str]] = Field(default=None, description="Optional base64 image strings for vision models")
    file_ids: Optional[List[str]] = Field(default=None, description="IDs of previously uploaded files")

    @model_validator(mode="before")
    @classmethod
    def handle_camel_case(cls, values: Any) -> Any:
        if isinstance(values, dict):
            if "conversationId" in values and not values.get("conversation_id"):
                values["conversation_id"] = values["conversationId"]
            if "userId" in values and not values.get("user_id"):
                values["user_id"] = values["userId"]
            if "requestId" in values and not values.get("request_id"):
                values["request_id"] = values["requestId"]
            if "systemPrompt" in values and not values.get("system_prompt"):
                values["system_prompt"] = values["systemPrompt"]
            if "maxTokens" in values and not values.get("max_tokens"):
                values["max_tokens"] = values["maxTokens"]
            if "fileIds" in values and not values.get("file_ids"):
                values["file_ids"] = values["fileIds"]
            if "taskType" in values and not values.get("task_type"):
                values["task_type"] = values["taskType"]
        return values


def validate_and_sanitize_image(image_str: str) -> str:
    """Validate image string (data URI or raw base64). Returns clean image string or raises HTTPException 400."""
    if not image_str or not isinstance(image_str, str):
        raise HTTPException(status_code=400, detail="Invalid image input: image data must be a non-empty string")
    
    raw_b64 = image_str
    if image_str.startswith("data:") and "," in image_str:
        raw_b64 = image_str.split(",", 1)[1]
    
    raw_b64 = raw_b64.strip()
    try:
        import base64
        import io
        from PIL import Image
        decoded = base64.b64decode(raw_b64, validate=True)
        if len(decoded) < 10:
            raise ValueError("Decoded image data too small")
        with Image.open(io.BytesIO(decoded)) as img:
            img.verify()
    except Exception as err:
        raise HTTPException(status_code=400, detail=f"Malformed image data: invalid base64 image ({err})")
    
    return image_str




# ================================================================
# FILE UPLOAD & EXTRACTION UTILITIES
# ================================================================

UPLOAD_DIR = Path(_ai_root) / "data" / "uploads"
UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
MAX_FILE_SIZE = 25 * 1024 * 1024
MAX_TOTAL_FILES_PER_REQUEST = 10

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"}
TEXT_EXTENSIONS = {
    ".txt", ".md", ".rtf", ".json", ".xml", ".yaml", ".yml", ".py", ".js", ".ts",
    ".jsx", ".tsx", ".java", ".c", ".cpp", ".h", ".hpp", ".cs", ".go", ".rs",
    ".php", ".sql", ".sh", ".ps1", ".ini", ".cfg", ".conf"
}

def sanitize_filename(filename: str) -> str:
    if not filename:
        return "unnamed_file"
    return Path(filename).name

def get_file_category(extension: str) -> str:
    extension = extension.lower()
    if extension in IMAGE_EXTENSIONS:
        return "image"
    if extension == ".pdf":
        return "pdf"
    if extension in {".doc", ".docx"}:
        return "document"
    if extension in {".xls", ".xlsx", ".csv"}:
        return "spreadsheet"
    if extension in {".ppt", ".pptx"}:
        return "presentation"
    if extension in TEXT_EXTENSIONS:
        return "text"
    return "binary"

@router.post("/chat/upload", summary="Upload files for AI processing")
@router.post("/api/chat/upload", summary="Upload files for AI processing")
async def upload_chat_files(files: List[UploadFile] = File(...)):
    if not files:
        raise HTTPException(status_code=400, detail="No files were uploaded.")
    if len(files) > MAX_TOTAL_FILES_PER_REQUEST:
        raise HTTPException(status_code=400, detail=f"Maximum {MAX_TOTAL_FILES_PER_REQUEST} files are allowed per request.")
    
    uploaded_files = []
    try:
        for upload in files:
            original_name = sanitize_filename(upload.filename or "")
            extension = Path(original_name).suffix.lower()
            file_id = f"{uuid.uuid4().hex}{extension}"
            destination = UPLOAD_DIR / file_id
            
            total_size = 0
            try:
                with destination.open("wb") as output:
                    while True:
                        chunk = await upload.read(1024 * 1024)
                        if not chunk:
                            break
                        total_size += len(chunk)
                        if total_size > MAX_FILE_SIZE:
                            destination.unlink(missing_ok=True)
                            raise HTTPException(status_code=413, detail=f"File '{original_name}' exceeds the 25 MB limit.")
                        output.write(chunk)
            finally:
                await upload.close()
                
            mime_type = upload.content_type or mimetypes.guess_type(original_name)[0] or "application/octet-stream"
            category = get_file_category(extension)
            
            file_info = {
                "file_id": file_id,
                "original_name": original_name,
                "stored_name": file_id,
                "path": str(destination),
                "mime_type": mime_type,
                "extension": extension,
                "category": category,
                "size": total_size,
                "status": "uploaded",
            }
            uploaded_files.append(file_info)
            logger.info(f"[UPLOAD] {original_name} -> {file_id} ({total_size} bytes)")
            
        return {"status": "success", "files": uploaded_files, "count": len(uploaded_files)}
    except HTTPException:
        raise
    except Exception as exc:
        logger.error(f"[UPLOAD] Upload failed: {exc}", exc_info=True)
        raise HTTPException(status_code=500, detail=f"File upload failed: {exc}")

def resolve_uploaded_file(file_id: str) -> Path:
    if not file_id:
        raise HTTPException(status_code=400, detail="Invalid file ID.")
    safe_name = Path(file_id).name
    if safe_name != file_id:
        raise HTTPException(status_code=400, detail="Invalid file ID.")
    path = UPLOAD_DIR / safe_name
    if not path.exists():
        raise HTTPException(status_code=404, detail=f"Uploaded file '{file_id}' was not found.")
    return path

async def extract_file_content(path: Path) -> Dict[str, Any]:
    extension = path.suffix.lower()
    result = {
        "file_name": path.name,
        "extension": extension,
        "category": get_file_category(extension),
        "text": "",
        "metadata": {},
        "success": True,
    }
    
    try:
        if extension in TEXT_EXTENSIONS or extension == ".csv":
            text = await asyncio.to_thread(path.read_text, encoding="utf-8", errors="ignore")
            result["text"] = text[:100_000]
            return result
            
        if extension == ".pdf":
            try:
                from pypdf import PdfReader
                def read_pdf():
                    reader = PdfReader(str(path))
                    return "\n".join([page.extract_text() or "" for page in reader.pages])
                text = await asyncio.to_thread(read_pdf)
                result["text"] = text[:100_000]
            except ImportError:
                result["success"] = False
                result["error"] = "pypdf is not installed."
            return result
            
        if extension == ".docx":
            try:
                from docx import Document
                def read_docx():
                    doc = Document(str(path))
                    return "\n".join([p.text for p in doc.paragraphs if p.text.strip()])
                text = await asyncio.to_thread(read_docx)
                result["text"] = text[:100_000]
            except ImportError:
                result["success"] = False
                result["error"] = "python-docx is not installed."
            return result
            
        if extension in {".xlsx", ".xls"}:
            try:
                import openpyxl
                def read_excel():
                    if extension == ".xlsx":
                        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
                        sheets = []
                        for sheet in wb.worksheets:
                            sheets.append(f"\n[SHEET: {sheet.title}]")
                            for row in sheet.iter_rows(values_only=True):
                                vals = [str(v) if v is not None else "" for v in row]
                                if any(vals): sheets.append(" | ".join(vals))
                        return "\n".join(sheets)
                    return "Legacy .xls files require an XLS-compatible reader."
                text = await asyncio.to_thread(read_excel)
                result["text"] = text[:100_000]
            except ImportError:
                result["success"] = False
                result["error"] = "openpyxl is not installed."
            return result
            
        if extension == ".pptx":
            try:
                from pptx import Presentation
                def read_pptx():
                    prs = Presentation(str(path))
                    slides = []
                    for i, slide in enumerate(prs.slides, 1):
                        slides.append(f"\n[SLIDE {i}]")
                        for shape in slide.shapes:
                            if hasattr(shape, "text") and shape.text.strip():
                                slides.append(shape.text.strip())
                    return "\n".join(slides)
                text = await asyncio.to_thread(read_pptx)
                result["text"] = text[:100_000]
            except ImportError:
                result["success"] = False
                result["error"] = "python-pptx is not installed."
            return result
            
        if extension in IMAGE_EXTENSIONS:
            result["metadata"] = {"requires_vision": True, "message": "Image should be processed through multimodal pipeline."}
            return result
            
        result["metadata"] = {"message": "File uploaded successfully but no text extractor is configured."}
        return result
        
    except Exception as exc:
        logger.warning(f"[FILE] Extraction failed for {path.name}: {exc}")
        result["success"] = False
        result["error"] = str(exc)
        return result

async def process_attached_files(file_ids: Optional[List[str]]) -> List[Dict[str, Any]]:
    if not file_ids: return []
    if len(file_ids) > MAX_TOTAL_FILES_PER_REQUEST:
        raise HTTPException(status_code=400, detail=f"Maximum {MAX_TOTAL_FILES_PER_REQUEST} files allowed.")
    
    processed = []
    for file_id in file_ids:
        path = resolve_uploaded_file(file_id)
        extracted = await extract_file_content(path)
        extracted["file_id"] = file_id
        extracted["path"] = str(path)
        processed.append(extracted)
    return processed

def build_file_context(files: List[Dict[str, Any]]) -> str:
    if not files: return ""
    sections = []
    for f in files:
        name = f.get("file_name", "unknown")
        cat = f.get("category", "unknown")
        succ = f.get("success", True)
        sec = f"\n===== FILE: {name} =====\nTYPE: {cat}\nEXTRACTION_SUCCESS: {succ}\n"
        if f.get("text"):
            sec += f"\nEXTRACTED CONTENT:\n{f['text']}\n"
        elif f.get("error"):
            sec += f"\nEXTRACTION ERROR:\n{f['error']}\n"
        sections.append(sec)
    return "\n".join(sections)


class ChatResponse(BaseModel):
    """Chat response payload."""
    response: str = Field(..., description="Generated text response")
    model_used: str = Field(..., description="Model used for generation")
    task_type: Optional[str] = Field(None, description="Classified task type")
    model_routing_reason: Optional[str] = Field(None, description="Reason model was chosen")
    status: str = Field(default="completed", description="Status of the request")
    agent_used: Optional[str] = Field(default="general", description="Agent used")
    task_id: Optional[str] = Field(None, description="Orchestrator task ID if agentic workflow was run")
    generated_files: Optional[List[Dict[str, Any]]] = Field(default_factory=list, description="Files generated by agent tools")
    tool_calls: Optional[List[Dict[str, Any]]] = Field(default_factory=list, description="Tool calls executed during run")
    tool_executions: Optional[List[Dict[str, Any]]] = Field(default_factory=list, description="Detailed tool executions with stdout and latency")
    ocr: Optional[Dict[str, Any]] = Field(None, description="Exact OCR extraction details")
    vision: Optional[Dict[str, Any]] = Field(None, description="Visual understanding details")
    combined_context: Optional[str] = Field(None, description="Unified OCR + Vision context")
    files: Optional[List[Dict[str, Any]]] = Field(default_factory=list, description="Extracted file contexts attached to the request")


@router.post("/chat", response_model=ChatResponse, summary="Chat with Sovereign AI")
@router.post("/api/chat", response_model=ChatResponse, summary="Chat with Sovereign AI")
async def chat(request: ChatRequest):
    """
    Send a message and get a complete response.
    When images are attached:
    1. Dedicated OCR (PaddleOCR) extracts exact alphanumeric text facts.
    2. Vision Service (Moondream) extracts visual scene understanding.
    3. Outputs are unified into a combined context.
    4. Downstream Agent / LLM reasons over the combined context.
    When no images are attached:
    Regular text-only chat without OCR/Vision invocation.
    """
    try:
        request_id = request.request_id or uuid.uuid4().hex
        logger.info(f"[{request_id}] CHAT_REQUEST conversation={request.conversation_id} user={request.user_id}")
        query = request.message.strip()
        has_images = bool(request.images and len(request.images) > 0)

        # -----------------------------------------------------------------
        # MULTIMODAL FLOW: Image attached
        # -----------------------------------------------------------------
        if has_images:
            task_type = "multimodal_inspection"
            primary_image = validate_and_sanitize_image(request.images[0])
            logger.info(f"[CHAT] Multimodal image request received: '{query[:60]}'...")

            # 1. Run Multimodal Orchestration (OCR + Moondream)
            multimodal_result = await multimodal_orchestrator.analyze(
                image_source=primary_image,
                user_prompt=query
            )

            ocr_data = multimodal_result.get("ocr", {})
            vision_data = multimodal_result.get("vision", {})
            unified_context = multimodal_result.get("combined_context", "")

            # 2. Select reasoning model (Chat/Reasoning LLM, not Moondream directly)
            reasoning_model = (
                request.model
                if request.model != "auto"
                else model_router.get_roles().get("chat", settings.OLLAMA_CHAT_MODEL)
            )

            # Retrieve Memory Context (STM, LTM, Mem0, Qdrant, Neo4j)
            memory_ctx = ""
            try:
                memory_ctx = await memory_manager.get_context(
                    query=query,
                    user_id=request.user_id,
                    conversation_id=request.conversation_id,
                    agent_id=request.agent
                )
            except Exception as mem_err:
                logger.warning(f"Memory retrieval warning in multimodal chat: {mem_err}")

            agent_prompt = multimodal_orchestrator.format_agent_prompt(
                unified_context=unified_context,
                user_query=query
            )

            eff_system_prompt = request.system_prompt or "You are an industrial engineering and risk assessment expert."
            if memory_ctx:
                eff_system_prompt = f"{eff_system_prompt}\n\n{memory_ctx}".strip()
        attached_files = await process_attached_files(request.file_ids)
        file_context = build_file_context(attached_files)
        if file_context:
            eff_system_prompt = f"{eff_system_prompt}\n\n===== ATTACHED FILE CONTEXT =====\n{file_context}"


            # 3. Downstream Reasoning LLM execution
            messages = [
                {"role": "system", "content": eff_system_prompt},
                {"role": "user", "content": agent_prompt}
            ]

            options = {}
            if request.temperature is not None:
                options["temperature"] = request.temperature
            if request.max_tokens is not None:
                options["num_predict"] = request.max_tokens

            reply = await ollama_client.chat(model=reasoning_model, messages=messages, options=options if options else None)

            # Trigger asynchronous memory extraction in background with plain text facts
            if request.user_id and reply:
                extracted_plain_facts = query
                if ocr_data.get("text"):
                    extracted_plain_facts += f"\n[Extracted Text]: {ocr_data['text']}"
                if vision_data.get("description"):
                    extracted_plain_facts += f"\n[Visual Description]: {vision_data['description']}"

                memory_manager.process_interaction_async(
                    user_id=request.user_id,
                    user_message=extracted_plain_facts,
                    assistant_response=reply,
                    conversation_id=request.conversation_id
                )

            logger.info(
                f"[TURN SUMMARY] user_id={request.user_id} | conv_id={request.conversation_id} | "
                f"memory_injected={len(memory_ctx)} chars | multimodal=True | model={reasoning_model}"
            )

            return ChatResponse(
                response=reply,
                model_used=reasoning_model,
                task_type=task_type,
                model_routing_reason="Multimodal vision and OCR reasoning model",
                agent_used=request.agent or "risk_inspector",
                status="completed",
                ocr=ocr_data,
                vision=vision_data,
                combined_context=unified_context,
                files=attached_files
            )

        # -----------------------------------------------------------------
        # REGULAR TEXT CHAT: No images attached
        # -----------------------------------------------------------------
        routing = task_classifier.classify(query, has_images=False)
        task_type = routing.task_type
        selected_agent = routing.agent if not request.agent or request.agent.lower() == "auto" else request.agent
        selected_model = model_router.route(query, task_type=task_type) if request.model == "auto" else request.model
        routing_reason = f"task classifier selected {selected_agent} ({routing.confidence:.2f})"

        logger.info(f"[ROUTER] Selected model '{selected_model}' for task '{task_type}' (reason: {routing_reason})")

        # Multi-Layer Centralized Context Assembly
        enriched = await central_context_builder.build(
            conversation_id=request.conversation_id or "default",
            user_id=request.user_id,
            query=query,
            agent_id=selected_agent,
            system_prompt=request.system_prompt or system_prompt_for(selected_agent)
        )
        attached_files = await process_attached_files(request.file_ids)
        file_context = build_file_context(attached_files)
        if file_context:
            enriched.system_prompt = enriched.system_prompt + "\n\n===== ATTACHED FILES =====\n" + file_context


        is_agent_mode = bool(selected_agent and selected_agent.lower() not in ["none", "direct"])

        if is_agent_mode:
            logger.info(f"[TASK_CLASSIFICATION] type={task_type} agent={selected_agent} confidence={routing.confidence:.2f}")
            logger.info(f"Executing LangGraph agent '{selected_agent}' for text query: {query[:60]}... (conv={request.conversation_id})")
            agent = get_langgraph_agent(
                model=selected_model,
                system_prompt=enriched.system_prompt,
                temperature=request.temperature,
                max_tokens=request.max_tokens
            )
            result = await agent.execute(
                query=query,
                history_messages=enriched.history_messages,
                conversation_id=request.conversation_id
            )

            task_id = f"task_{uuid.uuid4().hex[:8]}"
            first_tool = result.get("tool_calls", [{}])[0].get("tool") if result.get("tool_calls") else None
            try:
                await memory_manager.record_task_graph(
                    task_id=task_id,
                    task_type=task_type,
                    tool_name=first_tool
                )
            except Exception as graph_err:
                logger.warning(f"Neo4j record warning: {graph_err}")

            # Authoritative MongoDB persistence for conversation turns (Phase 3)
            if request.conversation_id:
                try:
                    await stm_manager.record_message(
                        conversation_id=request.conversation_id,
                        role="user",
                        content=query,
                        user_id=request.user_id,
                    )
                    await stm_manager.record_message(
                        conversation_id=request.conversation_id,
                        role="assistant",
                        content=result.get("final_response", ""),
                        user_id=request.user_id,
                        agent_id=selected_agent,
                        model=result.get("selected_model"),
                        execution_id=task_id,
                    )
                except Exception as persist_err:
                    logger.warning(f"[STM] Message persistence warning: {persist_err}")

            # Trigger asynchronous memory extraction and summarization in background
            if request.user_id and result.get("final_response"):
                memory_manager.process_interaction_async(
                    user_id=request.user_id,
                    user_message=query,
                    assistant_response=result["final_response"],
                    conversation_id=request.conversation_id
                )
            if request.conversation_id:
                asyncio.create_task(conversation_summarizer.summarize_if_needed(request.conversation_id, user_id=request.user_id))

            logger.info(
                f"[TURN SUMMARY] user_id={request.user_id} | conv_id={request.conversation_id} | "
                f"artifacts={len(enriched.artifacts)} | history_turns={len(enriched.history_messages)} | model={result['selected_model']}"
            )

            return ChatResponse(
                response=result["final_response"],
                model_used=result["selected_model"],
                task_type=task_type,
                model_routing_reason=routing_reason,
                agent_used=selected_agent,
                task_id=task_id,
                status="completed",
                generated_files=result.get("generated_files", []),
                tool_calls=result.get("tool_calls", []),
                tool_executions=result.get("tool_executions", [])
            )

        # Direct plain text LLM with full context and conversation turns
        messages = [{"role": "system", "content": enriched.system_prompt}]
        for hm in enriched.history_messages:
            messages.append({"role": hm["role"], "content": hm["content"]})
        messages.append({"role": "user", "content": query})

        options = {}
        if request.temperature is not None:
            options["temperature"] = request.temperature
        if request.max_tokens is not None:
            options["num_predict"] = request.max_tokens

        reply = await ollama_client.chat(model=selected_model, messages=messages, options=options if options else None)

        # Authoritative MongoDB persistence for direct conversation turns (Phase 3)
        if request.conversation_id:
            try:
                await stm_manager.record_message(
                    conversation_id=request.conversation_id,
                    role="user",
                    content=query,
                    user_id=request.user_id,
                )
                await stm_manager.record_message(
                    conversation_id=request.conversation_id,
                    role="assistant",
                    content=reply,
                    user_id=request.user_id,
                    agent_id="direct",
                    model=selected_model,
                )
            except Exception as persist_err:
                logger.warning(f"[STM] Message persistence warning: {persist_err}")

        # Trigger asynchronous memory extraction in background
        if request.user_id and reply:
            memory_manager.process_interaction_async(
                user_id=request.user_id,
                user_message=query,
                assistant_response=reply,
                conversation_id=request.conversation_id
            )
        if request.conversation_id:
            asyncio.create_task(conversation_summarizer.summarize_if_needed(request.conversation_id, user_id=request.user_id))

        logger.info(
            f"[TURN SUMMARY] user_id={request.user_id} | conv_id={request.conversation_id} | "
            f"artifacts={len(enriched.artifacts)} | history_turns={len(enriched.history_messages)} | model={selected_model}"
        )

        return ChatResponse(
            response=reply,
            model_used=selected_model,
            task_type=task_type,
            model_routing_reason=routing_reason,
            agent_used="direct",
            status="completed"
        )

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Chat endpoint error: {e}", exc_info=True)
        raise HTTPException(status_code=502, detail=f"AI service error: {str(e)}")


@router.post("/chat/stream", summary="Stream chat response")
@router.post("/api/chat/stream", summary="Stream chat response")
async def chat_stream(request: ChatRequest):
    """
    Stream a response token-by-token.
    When images are present:
    Emits progressive SSE events for OCR, Vision, Combining, and Reasoning,
    then streams the reasoned assessment.
    """
    try:
        request_id = request.request_id or uuid.uuid4().hex
        logger.info(f"[{request_id}] CHAT_STREAM_REQUEST conversation={request.conversation_id} user={request.user_id}")
        query = request.message.strip()
        has_images = bool(request.images and len(request.images) > 0)

        # -----------------------------------------------------------------
        # MULTIMODAL STREAMING FLOW
        # -----------------------------------------------------------------
        if has_images:
            primary_image = validate_and_sanitize_image(request.images[0])
            reasoning_model = (
                request.model
                if request.model != "auto"
                else model_router.get_roles().get("chat", settings.OLLAMA_CHAT_MODEL)
            )

            # Retrieve Memory Context (STM, LTM, Mem0, Qdrant, Neo4j)
            memory_ctx = ""
            try:
                memory_ctx = await memory_manager.get_context(
                    query=query,
                    user_id=request.user_id,
                    conversation_id=request.conversation_id,
                    agent_id=request.agent
                )
            except Exception as mem_err:
                logger.warning(f"Memory retrieval warning in multimodal stream: {mem_err}")

            eff_system_prompt = request.system_prompt or "You are an industrial engineering and risk assessment expert."
            if memory_ctx:
                eff_system_prompt = f"{eff_system_prompt}\n\n{memory_ctx}".strip()

            async def generate_multimodal():
                try:
                    # 1. Start notification
                    yield f"data: {json.dumps({'event': 'vision:progress', 'stage': 'uploading', 'message': 'Uploading image for multimodal inspection...'}, ensure_ascii=False)}\n\n"
                    await asyncio.sleep(0.05)

                    # 2. Dedicated OCR extraction
                    yield f"data: {json.dumps({'event': 'vision:progress', 'stage': 'ocr', 'message': 'Running OCR (extracting exact text)...'}, ensure_ascii=False)}\n\n"
                    loop = asyncio.get_running_loop()
                    try:
                        ocr_result = await loop.run_in_executor(None, multimodal_orchestrator.ocr.extract_text, primary_image)
                    except Exception as ocr_e:
                        ocr_result = {"success": False, "text": "", "error": str(ocr_e)}

                    # 3. Dedicated Moondream scene understanding
                    yield f"data: {json.dumps({'event': 'vision:progress', 'stage': 'vision', 'message': 'Analyzing visual scene (Moondream)...'}, ensure_ascii=False)}\n\n"
                    try:
                        vision_result = await multimodal_orchestrator.vision.analyze_visual_scene(primary_image)
                    except Exception as vis_e:
                        vision_result = {"success": False, "description": "Visual analysis unavailable.", "error": str(vis_e)}

                    # 4. Combining results
                    yield f"data: {json.dumps({'event': 'vision:progress', 'stage': 'combining', 'message': 'Combining facts and visual context...'}, ensure_ascii=False)}\n\n"
                    unified_context = multimodal_orchestrator.build_unified_context(
                        ocr_text=ocr_result.get("text", ""),
                        vision_description=vision_result.get("description", ""),
                        user_query=query
                    )

                    # Emit combined result event with metadata for UI
                    yield f"data: {json.dumps({'event': 'vision:result', 'ocr': ocr_result, 'vision': vision_result, 'combined_context': unified_context}, ensure_ascii=False)}\n\n"

                    # 5. Generating final reasoning response
                    yield f"data: {json.dumps({'event': 'vision:progress', 'stage': 'reasoning', 'message': 'Generating final risk analysis...'}, ensure_ascii=False)}\n\n"

                    agent_prompt = multimodal_orchestrator.format_agent_prompt(
                        unified_context=unified_context,
                        user_query=query
                    )

                    messages = [
                        {"role": "system", "content": eff_system_prompt},
                        {"role": "user", "content": agent_prompt}
                    ]

                    stream_options = {}
                    if request.temperature is not None:
                        stream_options["temperature"] = request.temperature
                    if request.max_tokens is not None:
                        stream_options["num_predict"] = request.max_tokens

                    streamed_tokens = []
                    async for token in ollama_client.chat_stream(
                        model=reasoning_model,
                        messages=messages,
                        options=stream_options if stream_options else None
                    ):
                        streamed_tokens.append(token)
                        payload = {"token": token, "content": token}
                        yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"

                    # Trigger memory extraction after multimodal streaming completes
                    full_reply = "".join(streamed_tokens)
                    if request.user_id and full_reply:
                        extracted_plain_facts = query
                        if ocr_result.get("text"):
                            extracted_plain_facts += f"\n[Extracted Text]: {ocr_result['text']}"
                        if vision_result.get("description"):
                            extracted_plain_facts += f"\n[Visual Description]: {vision_result['description']}"

                        memory_manager.process_interaction_async(
                            user_id=request.user_id,
                            user_message=extracted_plain_facts,
                            assistant_response=full_reply,
                            conversation_id=request.conversation_id
                        )

                    logger.info(
                        f"[TURN SUMMARY] user_id={request.user_id} | conv_id={request.conversation_id} | "
                        f"memory_injected={len(memory_ctx)} chars | multimodal=True | model={reasoning_model}"
                    )

                    # Send completion marker with multimodal metadata
                    yield f"data: {json.dumps({'status': 'completed', 'ocr': ocr_result, 'vision': vision_result}, ensure_ascii=False)}\n\n"
                    logger.info("[CHAT] Multimodal streaming completed successfully.")

                except Exception as e:
                    logger.error(f"Multimodal streaming error: {e}", exc_info=True)
                    yield f"data: {json.dumps({'error': str(e)}, ensure_ascii=False)}\n\n"

            return StreamingResponse(generate_multimodal(), media_type="text/event-stream")

        # -----------------------------------------------------------------
        # REGULAR TEXT STREAMING FLOW (NO OCR OR VISION)
        # -----------------------------------------------------------------
        routing = task_classifier.classify(query, has_images=False)
        task_type = routing.task_type
        selected_agent = routing.agent if not request.agent or request.agent.lower() == "auto" else request.agent
        selected_model = model_router.route(query, task_type=task_type) if request.model == "auto" else request.model
        routing_reason = f"task classifier selected {selected_agent} ({routing.confidence:.2f})"

        logger.info(
            f"[CHAT_STREAM] Query: '{query[:60]}' | user_id={request.user_id} | "
            f"conversation_id={request.conversation_id} | agent={selected_agent} | model={selected_model} (reason: {routing_reason})"
        )

        # Multi-Layer Centralized Context Assembly
        enriched = await central_context_builder.build(
            conversation_id=request.conversation_id or "default",
            user_id=request.user_id,
            query=query,
            agent_id=selected_agent,
            system_prompt=request.system_prompt or system_prompt_for(selected_agent)
        )

        is_agent_mode = bool(selected_agent and selected_agent.lower() not in ["none", "direct"])

        async def generate_text():
            try:
                yield f"data: {json.dumps({'event': 'chat:status', 'model': selected_model, 'task_type': task_type, 'agent': selected_agent, 'classification': routing.to_dict(), 'reason': routing_reason, 'status': 'started', 'intent': enriched.intent.value, 'sources': enriched.observability.get('rag_chunks_count', 0)}, ensure_ascii=False)}\n\n"
                streamed_tokens = []

                if is_agent_mode:
                    task_id = f"task_{uuid.uuid4().hex[:8]}"
                    try:
                        await memory_manager.record_task_graph(
                            task_id=task_id,
                            task_type=task_type
                        )
                    except Exception as graph_err:
                        logger.warning(f"Neo4j record warning: {graph_err}")

                    agent = get_langgraph_agent(
                        model=selected_model,
                        system_prompt=enriched.system_prompt,
                        temperature=request.temperature,
                        max_tokens=request.max_tokens
                    )
                    agent_full_resp = ""
                    async for event in agent.stream(
                        query=query,
                        history_messages=enriched.history_messages,
                        conversation_id=request.conversation_id
                    ):
                        if event.get("token"):
                            streamed_tokens.append(event["token"])
                        if event.get("content"):
                            streamed_tokens.append(event["content"])
                        if event.get("final_response"):
                            agent_full_resp = event["final_response"]
                        yield f"data: {json.dumps(event, ensure_ascii=False)}\n\n"

                    # Trigger memory extraction & rolling summarization
                    full_text = agent_full_resp or "".join(streamed_tokens)
                    if request.user_id and full_text:
                        memory_manager.process_interaction_async(
                            user_id=request.user_id,
                            user_message=query,
                            assistant_response=full_text,
                            conversation_id=request.conversation_id
                        )
                    if request.conversation_id:
                        asyncio.create_task(conversation_summarizer.summarize_if_needed(request.conversation_id, user_id=request.user_id))
                    yield f"data: {json.dumps({'status': 'completed', 'response': full_text, 'agent': selected_agent, 'task_type': task_type}, ensure_ascii=False)}\n\n"
                else:
                    messages = [{"role": "system", "content": enriched.system_prompt}]
                    for hm in enriched.history_messages:
                        messages.append({"role": hm["role"], "content": hm["content"]})
                    messages.append({"role": "user", "content": query})

                    stream_options = {}
                    if request.temperature is not None:
                        stream_options["temperature"] = request.temperature
                    if request.max_tokens is not None:
                        stream_options["num_predict"] = request.max_tokens

                    async for token in ollama_client.chat_stream(
                        model=selected_model,
                        messages=messages,
                        options=stream_options if stream_options else None
                    ):
                        streamed_tokens.append(token)
                        payload = {"token": token, "content": token}
                        yield f"data: {json.dumps(payload, ensure_ascii=False)}\n\n"

                    yield f"data: {json.dumps({'status': 'completed'}, ensure_ascii=False)}\n\n"

                    # Trigger memory extraction & rolling summarization
                    full_reply = "".join(streamed_tokens)
                    if request.user_id and full_reply:
                        memory_manager.process_interaction_async(
                            user_id=request.user_id,
                            user_message=query,
                            assistant_response=full_reply,
                            conversation_id=request.conversation_id
                        )
                    if request.conversation_id:
                        asyncio.create_task(conversation_summarizer.summarize_if_needed(request.conversation_id, user_id=request.user_id))

                logger.info(
                    f"[TURN SUMMARY] user_id={request.user_id} | conv_id={request.conversation_id} | "
                    f"artifacts={len(enriched.artifacts)} | history_turns={len(enriched.history_messages)} | model={selected_model}"
                )
                logger.info("Text chat streaming completed successfully.")

            except Exception as e:
                logger.error(f"Streaming generator error: {e}", exc_info=True)
                yield f"data: {json.dumps({'error': str(e)}, ensure_ascii=False)}\n\n"

        return StreamingResponse(generate_text(), media_type="text/event-stream")

    except HTTPException:
        raise
    except Exception as e:
        logger.error(f"Stream endpoint initialization error: {e}", exc_info=True)
        raise HTTPException(status_code=502, detail=f"Streaming error: {str(e)}")
