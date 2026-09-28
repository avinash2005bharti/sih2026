"""
Spreadsheet and tabular data tools for Sovereign AI Workbench.
Inspect, summarize, query, and write tabular CSV datasets.
"""

import csv
from typing import Any, Dict, List, Optional
from tools.base_tool import BaseTool
from tools.file_tool import _resolve_safe_path
from core.logging import logger


class SpreadsheetInspectTool(BaseTool):
    """Inspects a CSV file, computing row counts, columns, types, and numeric stats."""

    name = "inspect_spreadsheet"
    description = (
        "Inspect a CSV spreadsheet to get column names, total row count, inferred column types, "
        "basic summary statistics (min, max, average) for numeric columns, and preview sample rows."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "Relative path to CSV file inside workspace (e.g. 'telemetry.csv')"
            },
            "sample_rows": {
                "type": "integer",
                "description": "Number of sample rows to preview (default 5)"
            }
        },
        "required": ["file_path"]
    }

    async def arun(self, file_path: str, sample_rows: int = 5, **kwargs) -> Dict[str, Any]:
        try:
            target = _resolve_safe_path(file_path)
            if not target.exists():
                return {"success": False, "error": f"CSV file not found: {file_path}"}

            with open(target, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.DictReader(f)
                columns = reader.fieldnames or []
                rows = list(reader)

            total_rows = len(rows)
            preview_count = max(1, min(sample_rows, 20))
            sample = rows[:preview_count]

            col_stats = {}
            for col in columns:
                vals = [r[col].strip() for r in rows if r.get(col) is not None and r[col].strip() != ""]
                numeric_vals = []
                for v in vals:
                    try:
                        numeric_vals.append(float(v))
                    except ValueError:
                        pass

                if len(numeric_vals) == len(vals) and len(numeric_vals) > 0:
                    col_stats[col] = {
                        "type": "numeric",
                        "min": min(numeric_vals),
                        "max": max(numeric_vals),
                        "avg": round(sum(numeric_vals) / len(numeric_vals), 3),
                        "count": len(numeric_vals)
                    }
                else:
                    col_stats[col] = {
                        "type": "text/categorical",
                        "unique_count": len(set(vals)),
                        "count": len(vals)
                    }

            return {
                "success": True,
                "file_path": file_path,
                "total_rows": total_rows,
                "column_count": len(columns),
                "columns": columns,
                "column_analysis": col_stats,
                "sample_preview": sample
            }
        except Exception as e:
            logger.error(f"SpreadsheetInspectTool error on {file_path}: {e}")
            return {"success": False, "error": str(e)}


class SpreadsheetFilterTool(BaseTool):
    """Filters rows from a CSV based on column matching conditions."""

    name = "filter_spreadsheet"
    description = (
        "Filter rows from a CSV file matching a specific column condition "
        "(e.g. column 'status' equals 'FAIL', or numeric 'vibration' greater than 4.5)."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "Relative path to CSV file"
            },
            "column": {
                "type": "string",
                "description": "Column name to filter on"
            },
            "operator": {
                "type": "string",
                "enum": ["==", "!=", ">", "<", ">=", "<=", "contains"],
                "description": "Comparison operator (default '==')"
            },
            "value": {
                "type": "string",
                "description": "Value to compare against"
            },
            "limit": {
                "type": "integer",
                "description": "Maximum matching rows to return (default 20)"
            }
        },
        "required": ["file_path", "column", "value"]
    }

    async def arun(self, file_path: str, column: str, value: str, operator: str = "==", limit: int = 20, **kwargs) -> Dict[str, Any]:
        try:
            target = _resolve_safe_path(file_path)
            if not target.exists():
                return {"success": False, "error": f"CSV file not found: {file_path}"}

            with open(target, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.DictReader(f)
                rows = list(reader)

            matched = []
            max_results = max(1, min(limit, 100))

            for row in rows:
                cell = row.get(column, "")
                is_match = False

                if operator == "contains":
                    is_match = str(value).lower() in str(cell).lower()
                else:
                    try:
                        n_cell = float(cell)
                        n_val = float(value)
                        if operator == "==": is_match = (n_cell == n_val)
                        elif operator == "!=": is_match = (n_cell != n_val)
                        elif operator == ">": is_match = (n_cell > n_val)
                        elif operator == "<": is_match = (n_cell < n_val)
                        elif operator == ">=": is_match = (n_cell >= n_val)
                        elif operator == "<=": is_match = (n_cell <= n_val)
                    except ValueError:
                        if operator == "==": is_match = (str(cell) == str(value))
                        elif operator == "!=": is_match = (str(cell) != str(value))

                if is_match:
                    matched.append(row)
                    if len(matched) >= max_results:
                        break

            return {
                "success": True,
                "file_path": file_path,
                "column": column,
                "operator": operator,
                "filter_value": value,
                "matches_found": len(matched),
                "matches_count": len(matched),
                "rows": matched,
                "matched_rows": matched
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


class SpreadsheetWriteTool(BaseTool):
    """Creates a new CSV file or appends rows to an existing spreadsheet."""

    name = "write_spreadsheet"
    description = (
        "Create a new CSV spreadsheet with specified columns and rows, or append rows to an existing CSV."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "Relative path to target CSV file"
            },
            "columns": {
                "type": "array",
                "items": {"type": "string"},
                "description": "List of header column names"
            },
            "rows": {
                "type": "array",
                "items": {
                    "type": "object",
                    "description": "Row dictionary mapping column names to values"
                },
                "description": "List of rows to write"
            },
            "append": {
                "type": "boolean",
                "description": "If true, append to existing CSV instead of overwriting"
            }
        },
        "required": ["file_path", "columns", "rows"]
    }

    async def arun(self, file_path: str, columns: List[str], rows: List[Dict[str, Any]], append: bool = False, **kwargs) -> Dict[str, Any]:
        try:
            target = _resolve_safe_path(file_path)
            target.parent.mkdir(parents=True, exist_ok=True)

            file_exists = target.exists()
            mode = "a" if append and file_exists else "w"

            with open(target, mode, newline="", encoding="utf-8") as f:
                writer = csv.DictWriter(f, fieldnames=columns)
                if not append or not file_exists:
                    writer.writeheader()
                for r in rows:
                    writer.writerow(r)

            return {
                "success": True,
                "file_path": file_path,
                "rows_written": len(rows),
                "mode": "appended" if append and file_exists else "created"
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


inspect_spreadsheet_tool = SpreadsheetInspectTool()
filter_spreadsheet_tool = SpreadsheetFilterTool()
write_spreadsheet_tool = SpreadsheetWriteTool()


class SpreadsheetCreatorTool(BaseTool):
    """Generates formatted Excel (.xlsx) workbooks with headers, data, formulas, and auto-styling."""

    name = "xlsx_creator"
    description = "Generate a formatted Excel workbook (.xlsx) with headers, rows, and optional summary formulas."

    async def execute(
        self,
        title: str,
        headers: List[str],
        rows: List[List[Any]],
        filename: Optional[str] = None,
        summary_formula: Optional[str] = None,
        sheet_name: Optional[str] = None,
        **kwargs
    ) -> Dict[str, Any]:
        return await self.arun(
            title=title,
            headers=headers,
            rows=rows,
            filename=filename,
            summary_formula=summary_formula,
            sheet_name=sheet_name,
            **kwargs
        )

    async def arun(
        self,
        title: str,
        headers: List[str],
        rows: List[List[Any]],
        filename: Optional[str] = None,
        summary_formula: Optional[str] = None,
        sheet_name: Optional[str] = None,
        **kwargs
    ) -> Dict[str, Any]:
        try:
            import time
            import openpyxl
            from openpyxl.styles import Font, PatternFill, Alignment
            from core.config import settings
            from core.security import validate_safe_path

            fname = filename or f"spreadsheet_{int(time.time())}.xlsx"
            if not fname.lower().endswith(".xlsx"):
                fname += ".xlsx"

            target_path = settings.ARTIFACT_DIR / fname
            target_path.parent.mkdir(parents=True, exist_ok=True)
            safe_path = validate_safe_path(str(target_path))

            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = (sheet_name or title or "Sheet1")[:31]

            # Write header row
            header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
            header_font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
            ws.append(headers)
            for col_idx in range(1, len(headers) + 1):
                cell = ws.cell(row=1, column=col_idx)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(horizontal="center", vertical="center")

            # Write data rows
            data_font = Font(name="Arial", size=10)
            for r_data in rows:
                ws.append(r_data)
                curr_row = ws.max_row
                for col_idx in range(1, len(r_data) + 1):
                    ws.cell(row=curr_row, column=col_idx).font = data_font

            # Write summary formula row if provided
            if summary_formula:
                summary_row = ["Summary / Total"] + [""] * max(0, len(headers) - 2) + [summary_formula]
                ws.append(summary_row)
                s_row = ws.max_row
                summary_font = Font(name="Arial", size=10, bold=True)
                for col_idx in range(1, len(summary_row) + 1):
                    ws.cell(row=s_row, column=col_idx).font = summary_font

            # Auto-adjust column widths
            for col in ws.columns:
                max_len = max(len(str(cell.value or '')) for cell in col)
                col_letter = openpyxl.utils.get_column_letter(col[0].column)
                ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

            wb.save(str(safe_path))
            return {
                "status": "success",
                "success": True,
                "file_path": str(safe_path),
                "filename": fname,
                "row_count": ws.max_row,
                "col_count": ws.max_column
            }
        except Exception as e:
            logger.error(f"SpreadsheetCreatorTool error: {e}")
            return {"status": "error", "success": False, "error": str(e)}


spreadsheet_creator_tool = SpreadsheetCreatorTool()


