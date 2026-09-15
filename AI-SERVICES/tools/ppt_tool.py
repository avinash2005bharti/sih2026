"""
PowerPoint (.pptx) generation tool for Sovereign AI Workbench.
Creates structured, templated presentations using python-pptx in the workspace sandbox.
"""

from typing import Any, Dict, List, Optional
from pathlib import Path
from tools.base_tool import BaseTool
from tools.agent_tools import resolve_safe_path, SANDBOX_DIR
from memory.artifacts.artifact_store import artifact_store
from core.logging import logger

class CreatePPTTool(BaseTool):
    """Generates PowerPoint presentations (.pptx) inside the workspace sandbox."""

    name = "create_ppt"
    description = (
        "Create an official Microsoft PowerPoint (.pptx) presentation inside the workspace sandbox. "
        "Supports creating title slides and bulleted content slides. "
        "Useful for generating executive summaries, briefings, and automated reports."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_name": {
                "type": "string",
                "description": "Name of the output PPT file (e.g. 'Executive_Summary.pptx')"
            },
            "slides": {
                "type": "array",
                "items": {
                    "type": "object",
                    "properties": {
                        "type": {
                            "type": "string",
                            "enum": ["title", "content"],
                            "description": "Type of slide: 'title' for main title slide, 'content' for text/bullet points"
                        },
                        "title": {
                            "type": "string",
                            "description": "Title of the slide"
                        },
                        "content": {
                            "type": "string",
                            "description": "For 'title' slides, this is the subtitle. For 'content' slides, this is the main text body."
                        }
                    },
                    "required": ["type", "title"]
                },
                "description": "List of slides to include in the presentation"
            }
        },
        "required": ["file_name", "slides"]
    }

    async def arun(
        self,
        file_name: str,
        slides: List[Dict[str, Any]],
        **kwargs
    ) -> Dict[str, Any]:
        if not file_name.lower().endswith(".pptx"):
            file_name += ".pptx"

        try:
            target = resolve_safe_path(file_name)
            target.parent.mkdir(parents=True, exist_ok=True)

            try:
                from pptx import Presentation
                from pptx.util import Inches

                prs = Presentation()

                for slide_data in slides:
                    slide_type = slide_data.get("type", "content")
                    title_text = slide_data.get("title", "")
                    content_text = slide_data.get("content", "")

                    if slide_type == "title":
                        # Title Slide Layout (usually layout 0)
                        title_slide_layout = prs.slide_layouts[0]
                        slide = prs.slides.add_slide(title_slide_layout)
                        title = slide.shapes.title
                        subtitle = slide.placeholders[1]

                        title.text = title_text
                        subtitle.text = content_text

                    elif slide_type == "content":
                        # Title and Content Layout (usually layout 1)
                        bullet_slide_layout = prs.slide_layouts[1]
                        slide = prs.slides.add_slide(bullet_slide_layout)
                        shapes = slide.shapes

                        title_shape = shapes.title
                        body_shape = shapes.placeholders[1]

                        title_shape.text = title_text
                        tf = body_shape.text_frame
                        tf.text = content_text
                    else:
                        # Fallback for unknown type
                        blank_slide_layout = prs.slide_layouts[6]
                        slide = prs.slides.add_slide(blank_slide_layout)
                        left = top = width = height = Inches(1)
                        txBox = slide.shapes.add_textbox(left, top, width, height)
                        tf = txBox.text_frame
                        tf.text = title_text

                prs.save(str(target))

            except ImportError:
                return {"success": False, "error": "The 'python-pptx' library is required but not installed."}

            if not target.exists():
                return {"success": False, "error": f"Failed to generate PPT file '{file_name}' on disk."}

            size = target.stat().st_size
            rel_path = str(target.relative_to(SANDBOX_DIR)).replace("\\", "/")
            logger.info(f"[CreatePPTTool] Generated '{file_name}' ({size} bytes) at {rel_path}")

            # Register artifact in MongoDB if conversation_id provided
            conv_id = kwargs.get("conversation_id")
            if conv_id:
                artifact_store.register_artifact(
                    conversation_id=conv_id,
                    filename=target.name,
                    file_path=rel_path,
                    artifact_type="presentation",
                    mime_type="application/vnd.openxmlformats-officedocument.presentationml.presentation",
                    size_bytes=size,
                    description=f"PowerPoint presentation: {target.name}",
                    user_id=kwargs.get("user_id"),
                    execution_id=kwargs.get("execution_id")
                )

            return {
                "success": True,
                "file_name": target.name,
                "file_path": rel_path,
                "size_bytes": size,
                "mime_type": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                "message": f"PowerPoint presentation '{target.name}' generated successfully ({size} bytes)."
            }

        except Exception as e:
            logger.error(f"[CreatePPTTool] Error: {e}", exc_info=True)
            return {"success": False, "error": str(e)}

ppt_tool = CreatePPTTool()
