"""
Dedicated Excel (.xlsx) tool for Sovereign AI Workbench.
Creates styled, verified Excel workbooks using openpyxl in the workspace sandbox.
"""

import os
from pathlib import Path
from typing import Any, Dict, List, Optional, Union
from core.logging import logger
from tools.base_tool import BaseTool
from tools.agent_tools import resolve_safe_path, SANDBOX_DIR
from memory.artifacts.artifact_store import artifact_store


class CreateExcelTool(BaseTool):
    """Generates official formatted Excel spreadsheets (.xlsx) inside the workspace sandbox."""

    name = "create_excel"
    description = (
        "Create an official Microsoft Excel (.xlsx) spreadsheet file inside the sovereign workspace sandbox. "
        "Supports structured tables with columns, rows, automatic styling, and header formatting. "
        "Useful for maintenance inspection reports, equipment logs, sensor audits, and data exports."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_name": {
                "type": "string",
                "description": "Name of the output Excel file (e.g. 'Maintenance_Inspection_Report.xlsx')"
            },
            "headers": {
                "type": "array",
                "items": {"type": "string"},
                "description": "List of column header names (e.g. ['Item ID', 'Component', 'Status', 'Risk Level', 'Notes'])"
            },
            "rows": {
                "type": "array",
                "items": {
                    "type": "array",
                    "items": {"type": "string"}
                },
                "description": "List of rows, where each row is an array of cell values matching the headers"
            },
            "title": {
                "type": "string",
                "description": "Optional title heading placed at row 1"
            },
            "sheet_name": {
                "type": "string",
                "description": "Name of the worksheet tab (default: 'Inspection Report')"
            }
        },
        "required": ["file_name", "headers", "rows"]
    }

    async def arun(
        self,
        file_name: str,
        headers: List[str],
        rows: List[Union[List[Any], Dict[str, Any]]],
        title: Optional[str] = None,
        sheet_name: str = "Inspection Report",
        **kwargs
    ) -> Dict[str, Any]:
        if not file_name.lower().endswith(".xlsx"):
            file_name += ".xlsx"

        try:
            target = resolve_safe_path(file_name)
            target.parent.mkdir(parents=True, exist_ok=True)

            try:
                import openpyxl
                from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
                from openpyxl.utils import get_column_letter

                wb = openpyxl.Workbook()
                ws = wb.active
                ws.title = (sheet_name or "Report")[:31]

                current_row = 1

                # Optional title block
                if title:
                    ws.cell(row=current_row, column=1, value=title)
                    ws.cell(row=current_row, column=1).font = Font(size=14, bold=True, color="1E293B")
                    current_row += 2

                # Header row styling
                header_fill = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid") # Navy Blue
                header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                header_alignment = Alignment(horizontal="center", vertical="center")
                thin_border = Border(
                    left=Side(style="thin", color="CBD5E1"),
                    right=Side(style="thin", color="CBD5E1"),
                    top=Side(style="thin", color="CBD5E1"),
                    bottom=Side(style="thin", color="CBD5E1")
                )

                for col_idx, h_text in enumerate(headers, 1):
                    cell = ws.cell(row=current_row, column=col_idx, value=str(h_text))
                    cell.fill = header_fill
                    cell.font = header_font
                    cell.alignment = header_alignment
                    cell.border = thin_border
                current_row += 1

                # Data rows styling
                alt_fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
                regular_font = Font(name="Calibri", size=10)

                for r_idx, row_data in enumerate(rows):
                    is_alt = (r_idx % 2 == 1)
                    if isinstance(row_data, dict):
                        row_vals = [row_data.get(h, "") for h in headers]
                    else:
                        row_vals = list(row_data)

                    for col_idx, val in enumerate(row_vals, 1):
                        cell = ws.cell(row=current_row, column=col_idx, value=val)
                        cell.font = regular_font
                        cell.border = thin_border
                        if is_alt:
                            cell.fill = alt_fill
                    current_row += 1

                # Auto-adjust column widths
                for col in ws.columns:
                    max_len = 0
                    col_letter = get_column_letter(col[0].column)
                    for cell in col:
                        val_str = str(cell.value or '')
                        max_len = max(max_len, len(val_str))
                    ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

                wb.save(str(target))

            except ImportError:
                return {"success": False, "error": "The 'openpyxl' library is required but not installed."}

            if not target.exists():
                return {"success": False, "error": f"Failed to generate Excel file '{file_name}' on disk."}

            size = target.stat().st_size
            rel_path = str(target.relative_to(SANDBOX_DIR)).replace("\\", "/")
            logger.info(f"[CreateExcelTool] Generated '{file_name}' ({size} bytes) at {rel_path}")

            # Register artifact in MongoDB if conversation_id provided
            conv_id = kwargs.get("conversation_id")
            if conv_id:
                artifact_store.register_artifact(
                    conversation_id=conv_id,
                    filename=target.name,
                    file_path=rel_path,
                    artifact_type="spreadsheet",
                    mime_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    size_bytes=size,
                    description=title or f"Excel spreadsheet: {target.name}",
                    user_id=kwargs.get("user_id"),
                    execution_id=kwargs.get("execution_id")
                )

            return {
                "success": True,
                "file_name": target.name,
                "file_path": rel_path,
                "size_bytes": size,
                "mime_type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "message": f"Excel spreadsheet '{target.name}' generated successfully ({size} bytes)."
            }

        except Exception as e:
            logger.error(f"[CreateExcelTool] Error: {e}", exc_info=True)
            return {"success": False, "error": str(e)}


excel_tool = CreateExcelTool()
