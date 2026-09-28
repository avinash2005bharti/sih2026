"""
Document analysis and generation tools for Sovereign AI Workbench.
Supports plain text, markdown, PDF generation (via ReportLab), and document inspection.
"""

import re
from pathlib import Path
from typing import Any, Dict, List, Optional
from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from tools.base_tool import BaseTool
from tools.file_tool import _resolve_safe_path, SANDBOX_DIR
from core.logging import logger

REPORTS_DIR = (SANDBOX_DIR / "reports").resolve()
REPORTS_DIR.mkdir(parents=True, exist_ok=True)


class GeneratePDFTool(BaseTool):
    """Generates official formatted PDF reports inside the sandbox reports directory."""

    name = "create_pdf"
    description = (
        "Generate a formatted PDF document inside the sovereign workspace reports directory. "
        "Useful for exporting formal summaries, audit reports, and technical manuals."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_name": {
                "type": "string",
                "description": "Name of the output PDF file (e.g. 'audit_report.pdf')"
            },
            "title": {
                "type": "string",
                "description": "Document title heading"
            },
            "content": {
                "type": "string",
                "description": "Main body text paragraphs (separated by double newlines)"
            }
        },
        "required": ["file_name", "title", "content"]
    }

    async def arun(self, file_name: str, title: str, content: str, **kwargs) -> Dict[str, Any]:
        if not file_name.endswith(".pdf"):
            file_name += ".pdf"

        clean_name = Path(file_name).name
        target = REPORTS_DIR / clean_name

        try:
            styles = getSampleStyleSheet()
            doc = SimpleDocTemplate(str(target), pagesize=letter)
            story = [
                Paragraph(title, styles['Title']),
                Spacer(1, 14),
            ]
            for para in content.split("\n\n"):
                if para.strip():
                    clean_p = para.strip().replace("\n", "<br/>")
                    story.append(Paragraph(clean_p, styles['Normal']))
                    story.append(Spacer(1, 10))

            doc.build(story)

            if not target.exists():
                return {"success": False, "error": f"Failed to create PDF '{clean_name}' on disk."}

            size = target.stat().st_size
            rel_path = f"reports/{clean_name}"
            logger.info(f"GeneratePDFTool generated '{clean_name}' ({size} bytes)")

            return {
                "success": True,
                "file_name": clean_name,
                "file_path": rel_path,
                "size_bytes": size,
                "mime_type": "application/pdf",
                "message": f"PDF '{clean_name}' generated successfully ({size} bytes)."
            }
        except Exception as e:
            logger.error(f"GeneratePDFTool error: {e}")
            return {"success": False, "error": str(e)}


class DocumentInspectTool(BaseTool):
    """Inspects text document statistics and structural overview."""

    name = "inspect_document"
    description = (
        "Inspect a document file to get metadata, line count, word count, character count, "
        "and preview the first and last few lines."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "Relative path to document inside workspace (e.g. 'manual.md', 'logs/system.log')"
            },
            "preview_lines": {
                "type": "integer",
                "description": "Number of preview lines from head and tail (default 10)"
            }
        },
        "required": ["file_path"]
    }

    async def arun(self, file_path: str, preview_lines: int = 10, **kwargs) -> Dict[str, Any]:
        try:
            target = _resolve_safe_path(file_path)
            if not target.exists():
                return {"success": False, "error": f"Document not found: {file_path}"}
            if not target.is_file():
                return {"success": False, "error": f"Path is not a regular file: {file_path}"}

            with open(target, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()

            lines = content.splitlines()
            words = content.split()
            preview_n = max(1, min(preview_lines, 50))

            return {
                "success": True,
                "file_path": file_path,
                "file_size_bytes": target.stat().st_size,
                "total_lines": len(lines),
                "total_words": len(words),
                "total_characters": len(content),
                "head_preview": lines[:preview_n],
                "tail_preview": lines[-preview_n:] if len(lines) > preview_n else []
            }
        except Exception as e:
            logger.error(f"DocumentInspectTool error on {file_path}: {e}")
            return {"success": False, "error": str(e)}


class DocumentExtractTool(BaseTool):
    """Searches and extracts sections or patterns from a document."""

    name = "extract_document_sections"
    description = (
        "Search a document for specific keywords, headings, or regex patterns and extract "
        "matching lines with surrounding context."
    )
    parameters = {
        "type": "object",
        "properties": {
            "file_path": {
                "type": "string",
                "description": "Relative path to the document"
            },
            "pattern": {
                "type": "string",
                "description": "Keyword or regular expression to search for"
            },
            "context_lines": {
                "type": "integer",
                "description": "Number of surrounding context lines to include (default 2)"
            },
            "max_matches": {
                "type": "integer",
                "description": "Maximum number of matching sections to return (default 10)"
            }
        },
        "required": ["file_path", "pattern"]
    }

    async def arun(self, file_path: str, pattern: str, context_lines: int = 2, max_matches: int = 10, **kwargs) -> Dict[str, Any]:
        try:
            target = _resolve_safe_path(file_path)
            if not target.exists():
                return {"success": False, "error": f"Document not found: {file_path}"}

            with open(target, "r", encoding="utf-8", errors="replace") as f:
                lines = f.readlines()

            regex = re.compile(pattern, re.IGNORECASE)
            matches = []
            c = max(0, min(context_lines, 10))
            limit = max(1, min(max_matches, 50))

            for idx, line in enumerate(lines):
                if regex.search(line):
                    start = max(0, idx - c)
                    end = min(len(lines), idx + c + 1)
                    matches.append({
                        "line_number": idx + 1,
                        "matched_line": line.strip(),
                        "context": [lines[i].strip() for i in range(start, end)]
                    })
                    if len(matches) >= limit:
                        break

            return {
                "success": True,
                "file_path": file_path,
                "pattern": pattern,
                "matches_found": len(matches),
                "matches": matches
            }
        except Exception as e:
            return {"success": False, "error": str(e)}


class DocumentCreatorTool(BaseTool):
    """Generates formatted Microsoft Word (.docx) documents with headings, paragraphs, and tables."""

    name = "docx_creator"
    description = "Generate a formatted Word document (.docx) with structured sections, headings, and tables."

    async def execute(
        self,
        title: str,
        sections: List[Dict[str, Any]],
        filename: Optional[str] = None,
        **kwargs
    ) -> Dict[str, Any]:
        return await self.arun(title=title, sections=sections, filename=filename, **kwargs)

    async def arun(
        self,
        title: str,
        sections: List[Dict[str, Any]],
        filename: Optional[str] = None,
        **kwargs
    ) -> Dict[str, Any]:
        try:
            import time
            import docx
            from core.config import settings
            from core.security import validate_safe_path

            fname = filename or f"document_{int(time.time())}.docx"
            if not fname.lower().endswith(".docx"):
                fname += ".docx"

            target_path = settings.ARTIFACT_DIR / fname
            target_path.parent.mkdir(parents=True, exist_ok=True)
            safe_path = validate_safe_path(str(target_path))

            doc = docx.Document()
            doc.add_heading(title, level=0)

            for sec in sections:
                h = sec.get("heading")
                if h:
                    doc.add_heading(h, level=1)
                content = sec.get("content") or sec.get("body")
                if content:
                    doc.add_paragraph(content)
                table_data = sec.get("table_data")
                if table_data and isinstance(table_data, list) and len(table_data) > 0:
                    num_rows = len(table_data)
                    num_cols = len(table_data[0]) if num_rows > 0 and isinstance(table_data[0], list) else 1
                    t = doc.add_table(rows=num_rows, cols=num_cols)
                    t.style = 'Table Grid'
                    for r_idx, row in enumerate(table_data):
                        if isinstance(row, list):
                            for c_idx, cell_val in enumerate(row):
                                if c_idx < num_cols:
                                    t.cell(r_idx, c_idx).text = str(cell_val)
                        else:
                            t.cell(r_idx, 0).text = str(row)

            doc.save(str(safe_path))
            return {
                "status": "success",
                "success": True,
                "file_path": str(safe_path),
                "filename": fname
            }
        except Exception as e:
            logger.error(f"DocumentCreatorTool error: {e}")
            return {"status": "error", "success": False, "error": str(e)}


document_creator_tool = DocumentCreatorTool()

